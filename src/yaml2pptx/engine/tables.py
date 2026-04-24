"""Table generation for python-pptx slides."""

from __future__ import annotations

from pptx.util import Inches

from yaml2pptx.models.elements import TableData


def add_table_to_slide(
    slide,  # type: ignore[no-untyped-def]
    table_data: TableData,
    left: float = 1.0,
    top: float = 2.0,
    width: float = 8.0,
    row_height: float = 0.4,
) -> None:
    """Add a table to a slide."""
    has_headers = table_data.headers is not None

    if not table_data.rows and not has_headers:
        return

    num_rows = len(table_data.rows) + (1 if has_headers else 0)

    if has_headers:
        num_cols = len(table_data.headers)  # type: ignore[arg-type]
    elif table_data.rows:
        num_cols = len(table_data.rows[0])
    else:
        return
    height = row_height * num_rows

    table_shape = slide.shapes.add_table(
        num_rows,
        num_cols,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    table = table_shape.table

    row_offset = 0
    if has_headers and table_data.headers:
        for col_idx, header in enumerate(table_data.headers):
            table.cell(0, col_idx).text = header
        row_offset = 1

    for row_idx, row in enumerate(table_data.rows):
        for col_idx, value in enumerate(row):
            table.cell(row_idx + row_offset, col_idx).text = str(value)
